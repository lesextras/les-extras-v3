"""Fichiers modifiables de la banque d'outils (Word, Excel, PowerPoint).

python3 office.py <specs.json> <sortie> [id]
Même mini-langage que le PDF (dsl.js). Les blocs HTML bruts sont ignorés :
un outil qui en porte fournit `officeCorps`.
"""
import html as H
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Emu, Pt as PPt

NAVY = "1E2746"
INDIGO = "5146D8"
INDIGO_L = "ECEAFD"
JAUNE_L = "FDF5D8"
GRIS = "5B6275"
LIGNE = "D9DCE6"
COULEURS = {"jaune": "F4C542", "indigo": INDIGO, "vert": "1F8A5B", "corail": "D9454B", "": NAVY}
PIED = "ADéPA · association loi 1901, Melun (77) · Outil gratuit, à utiliser et à copier librement dans votre structure, pas à revendre · les-extras.fr/ressources"
EMOJI_FONT = "Segoe UI Emoji"


def typo(t):
    t = re.sub(r" ([?!;:»])", " \\1", t)
    return t.replace("« ", "« ")


def morceaux(s):
    """HTML simple -> [(texte, gras, italique)]."""
    s = str(s or "")
    s = re.sub(r"<br\s*/?>", "\n", s)
    out = []
    gras = ita = False
    for part in re.split(r"(<[^>]+>)", s):
        if not part:
            continue
        if part.startswith("<"):
            tag = part.strip("</> ").split()[0].lower() if part.strip("</> ") else ""
            fermant = part.startswith("</")
            if tag in ("b", "strong"):
                gras = not fermant
            elif tag in ("i", "em"):
                ita = not fermant
            continue
        out.append((typo(H.unescape(part)), gras, ita))
    return out


def texte_brut(s):
    return "".join(t for t, _, _ in morceaux(s))


def ombrer(cell, couleur):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tcPr.append(shd)


def bordures(table, couleur=LIGNE, taille=6, style="single"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for bord in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{bord}")
        e.set(qn("w:val"), style)
        e.set(qn("w:sz"), str(taille))
        e.set(qn("w:color"), couleur)
        b.append(e)
    tblPr.append(b)


def ligne_basse(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "dotted")
    e.set(qn("w:sz"), "6")
    e.set(qn("w:color"), "9AA1B4")
    b.append(e)
    pPr.append(b)


def ecrire(p, s, taille=None, couleur=None, gras=None, police=None):
    for t, g, i in morceaux(s):
        r = p.add_run(t)
        r.bold = gras if gras is not None else g
        r.italic = i
        if taille:
            r.font.size = Pt(taille)
        if couleur:
            r.font.color.rgb = RGBColor.from_string(couleur)
        if police:
            r.font.name = police
    return p


def hauteur(row, mm):
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def titre_bloc(doc, titre, couleur=""):
    if not titre:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    ecrire(p, titre.upper(), 9.5, COULEURS.get(couleur, NAVY), True)


def largeur_utile(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def bloc_docx(doc, b):
    if isinstance(b, str):
        return
    t = b["t"]
    if t == "bandeau":
        tab = doc.add_table(rows=1, cols=1)
        c = tab.cell(0, 0)
        ombrer(c, INDIGO_L)
        ecrire(c.paragraphs[0], b["html"], 10)
        doc.add_paragraph()
    elif t == "texte":
        p = doc.add_paragraph()
        ecrire(p, b["html"], 8.5 if b.get("petit") else 10.5, GRIS if b.get("petit") else None)
    elif t == "tableau":
        titre_bloc(doc, b.get("titre"), b.get("couleur", ""))
        cols = b["cols"]
        lignes = b.get("lignes") or []
        n = 1 + len(lignes) + (b.get("vides") or 0) + (1 if b.get("total") else 0)
        tab = doc.add_table(rows=n, cols=len(cols))
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordures(tab)
        larg = largeur_utile(doc)
        poids = [c.get("w") or (100 / len(cols)) for c in cols]
        tot = sum(poids)
        tab.autofit = False
        for j, c in enumerate(cols):
            tab.columns[j].width = int(larg * poids[j] / tot)
            cell = tab.cell(0, j)
            ombrer(cell, INDIGO)
            ecrire(cell.paragraphs[0], c["t"], 9, "FFFFFF", True)
            for i in range(n):
                tab.cell(i, j).width = int(larg * poids[j] / tot)
        for i, l in enumerate(lignes):
            row = tab.rows[1 + i]
            if b.get("h"):
                hauteur(row, b["h"])
            for j, v in enumerate(l):
                p = row.cells[j].paragraphs[0]
                if j == 0 and b.get("emo"):
                    ecrire(p, v, 18, police=EMOJI_FONT)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    ecrire(p, v, 10, gras=True if (j == 0 and b.get("premiereGras")) else None)
        for k in range(b.get("vides") or 0):
            hauteur(tab.rows[1 + len(lignes) + k], b.get("h") or 9)
        if b.get("total"):
            row = tab.rows[-1]
            ecrire(row.cells[0].paragraphs[0], "Total", 10, gras=True)
            for c in row.cells:
                ombrer(c, "F5F6FB")
        doc.add_paragraph()
    elif t == "champs":
        titre_bloc(doc, b.get("titre"), b.get("couleur", ""))
        for it in b["items"]:
            lab, nb = (it if isinstance(it, list) else [it, 1])
            if lab:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(0)
                ecrire(p, lab, 8.5, INDIGO, True)
            for _ in range(nb):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(10)
                ligne_basse(p)
    elif t in ("puces", "aretenir"):
        titre_bloc(doc, b.get("titre"), b.get("couleur", "jaune" if t == "aretenir" else ""))
        for x in b["items"]:
            ecrire(doc.add_paragraph(style="List Bullet"), x, 10)
    elif t == "etapes":
        titre_bloc(doc, b.get("titre"))
        for tps, x in b["items"]:
            p = doc.add_paragraph(style="List Number")
            if tps:
                ecrire(p, tps + "  ", 10, INDIGO, True)
            ecrire(p, x, 10)
    elif t == "cartes":
        if b.get("titre"):
            titre_bloc(doc, b["titre"])
        items = list(b["items"]) + [["", ""]] * (b.get("vides") or 0)
        cols = b.get("cols") or 4
        rows = (len(items) + cols - 1) // cols
        tab = doc.add_table(rows=rows, cols=cols)
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordures(tab, "9AA1B4", 8, "dashed")
        for i, (e, m) in enumerate(items):
            c = tab.cell(i // cols, i % cols)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if e:
                ecrire(p, e, 30, police=EMOJI_FONT)
            p2 = c.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ecrire(p2, m, 12, NAVY, True)
        for row in tab.rows:
            hauteur(row, min(b.get("h") or 34, 60))
        doc.add_paragraph()
    elif t == "zone":
        tab = doc.add_table(rows=1, cols=1)
        bordures(tab, "9AA1B4", 8, "dashed")
        c = tab.cell(0, 0)
        ecrire(c.paragraphs[0], b["titre"].upper(), 8.5, INDIGO, True)
        if b.get("aide"):
            ecrire(c.add_paragraph(), b["aide"], 8.5, GRIS)
        hauteur(tab.rows[0], min(b.get("h") or 30, 200))
        doc.add_paragraph()
    elif t == "cocher":
        titre_bloc(doc, b.get("titre"), b.get("couleur", ""))
        for x in b["items"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            ecrire(p, "☐  ", 12, INDIGO)
            ecrire(p, x, 10)
    elif t == "echelle":
        if b.get("titre"):
            titre_bloc(doc, b["titre"])
        items = b["items"]
        tab = doc.add_table(rows=len(items), cols=2)
        bordures(tab)
        for i, it in enumerate(items):
            couleur, tt, aide = (list(it) + ["", "", ""])[:3]
            c0, c1 = tab.cell(i, 0), tab.cell(i, 1)
            c0.width = Mm(20)
            c1.width = largeur_utile(doc) - Mm(20)
            ombrer(c0, couleur.lstrip("#").upper())
            p = c0.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ecrire(p, str(len(items) - i), 22, "FFFFFF", True)
            ecrire(c1.paragraphs[0], tt, 13, NAVY, True)
            ecrire(c1.add_paragraph(), aide or "", 10)
            hauteur(tab.rows[i], 22)
        doc.add_paragraph()
    elif t == "cote":
        for x in b["blocs"]:
            bloc_docx(doc, x)
    elif t == "signatures":
        tab = doc.add_table(rows=1, cols=len(b["items"]))
        bordures(tab, "9AA1B4", 8, "dashed")
        for j, s in enumerate(b["items"]):
            ecrire(tab.cell(0, j).paragraphs[0], s.upper(), 8.5, INDIGO, True)
        hauteur(tab.rows[0], 24)
        doc.add_paragraph()
    elif t == "saut":
        doc.add_page_break()


def entete_docx(doc, r):
    p = doc.add_paragraph()
    ecrire(p, texte_brut(r["kicker"]).upper(), 8.5, INDIGO, True)
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    ecrire(p, texte_brut(r["t1"]).upper() + " ", 22, NAVY, True)
    ecrire(p, texte_brut(r["t2"]).upper(), 22, INDIGO, True)
    if r.get("sous"):
        p = doc.add_paragraph()
        ecrire(p, r["sous"], 12, INDIGO).runs[0].italic = True
    p = doc.add_paragraph()
    ecrire(p, f"{r['format']} · Gratuit · Association ADéPA", 8.5, GRIS)


def faire_docx(r, chemin):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(NAVY)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    if r.get("paysage") or r.get("diapos"):
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Mm(15))
    sec.top_margin, sec.bottom_margin = Mm(13), Mm(13)
    fp = sec.footer.paragraphs[0]
    mention = PIED + (" · Les exemples sont fictifs." if r.get("fictif") else "") + (" · Modèle à adapter et à faire valider par votre structure." if r.get("modele") else "")
    ecrire(fp, mention, 7.5, GRIS)
    entete_docx(doc, r)
    if r.get("officeCorps"):
        pages = [r["officeCorps"]]
    elif r.get("diapos"):
        pages = [[{"t": "puces", "titre": f"{i + 1}. {d[1]}", "items": d[2]}] for i, d in enumerate(r["diapos"])]
    else:
        pages = r.get("pages") or [r.get("corps") or []]
    for i, blocs in enumerate(pages):
        if i:
            doc.add_page_break()
        for b in blocs:
            bloc_docx(doc, b)
    doc.core_properties.author = "Association ADéPA"
    doc.core_properties.title = texte_brut(r["titre"])
    doc.save(chemin)


# ── Excel ────────────────────────────────────────────────────────────────
FIN = Side(style="thin", color=LIGNE)
CADRE = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)


def tableaux(r):
    sources = r.get("pages") or [r.get("corps") or []]
    out = []

    def fouiller(bs):
        for b in bs:
            if isinstance(b, str):
                continue
            if b["t"] == "tableau":
                out.append(b)
            elif b["t"] == "cote":
                fouiller(b["blocs"])

    for bs in sources:
        fouiller(bs)
    return out


def faire_xlsx(r, chemin):
    tabs = tableaux(r)
    if not tabs:
        raise SystemExit(f"{r['id']} : xlsx demandé sans tableau")
    wb = Workbook()
    wb.remove(wb.active)
    noms = set()
    for k, b in enumerate(tabs):
        nom = texte_brut(b.get("feuille") or b.get("titre") or (texte_brut(r["t1"]) + " " + texte_brut(r["t2"])))[:31]
        nom = re.sub(r"[\[\]\*\?/\\:]", "", nom).strip() or f"Tableau {k + 1}"
        while nom in noms:
            nom = (nom[:28] + f" {k + 1}")
        noms.add(nom)
        ws = wb.create_sheet(nom)
        cols = b["cols"]
        ws.cell(1, 1, texte_brut(b.get("titre") or (r["t1"] + " " + r["t2"])).upper()).font = Font(bold=True, size=14, color=NAVY)
        ws.cell(2, 1, texte_brut(r.get("sous") or r["description"])).font = Font(italic=True, size=10, color=GRIS)
        for j, c in enumerate(cols, 1):
            cell = ws.cell(3, j, texte_brut(c["t"]))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=INDIGO)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = CADRE
            w = c.get("w") or (100 / len(cols))
            ws.column_dimensions[get_column_letter(j)].width = max(8, round(w * (1.6 if r.get("paysage") else 1.05)))
        ws.row_dimensions[3].height = 22
        lignes = b.get("lignes") or []
        ligne = 4
        for l in lignes:
            for j, v in enumerate(l, 1):
                if isinstance(v, (int, float)):
                    val = v
                else:
                    val = texte_brut(v)
                    if not val.startswith("="):
                        val = val or None
                cell = ws.cell(ligne, j, val)
                cell.border = CADRE
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if j == 1 and b.get("premiereGras"):
                    cell.font = Font(bold=True)
            if b.get("h"):
                ws.row_dimensions[ligne].height = min(b["h"] * 2.6, 120)
            ligne += 1
        for _ in range(b.get("vides") or 0):
            for j in range(1, len(cols) + 1):
                ws.cell(ligne, j).border = CADRE
            ws.row_dimensions[ligne].height = min((b.get("h") or 9) * 2.6, 120)
            ligne += 1
        if b.get("ligneFormule"):
            col, modele = b["ligneFormule"]
            for rr in range(4, ligne):
                ws.cell(rr, col + 1, modele.replace("{r}", str(rr))).border = CADRE
        if b.get("total"):
            ws.cell(ligne, 1, "Total").font = Font(bold=True)
            for j in range(1, len(cols) + 1):
                cell = ws.cell(ligne, j)
                cell.border = CADRE
                cell.fill = PatternFill("solid", fgColor="F5F6FB")
            for idx in b["total"]:
                lettre = get_column_letter(idx + 1)
                cell = ws.cell(ligne, idx + 1, f"=SUM({lettre}4:{lettre}{ligne - 1})")
                cell.font = Font(bold=True)
            ligne += 1
        for idx in b.get("euros") or []:
            for rr in range(4, ligne):
                ws.cell(rr, idx + 1).number_format = '#,##0.00 "€"'
        ws.cell(ligne + 1, 1, PIED).font = Font(size=8, color=GRIS)
        ws.freeze_panes = "A4"
        ws.page_setup.orientation = "landscape" if r.get("paysage") else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "3:3"
    wb.properties.creator = "Association ADéPA"
    wb.save(chemin)


# ── PowerPoint ───────────────────────────────────────────────────────────
def boite(slide, x, y, w, h, texte, taille, couleur=NAVY, gras=False, italique=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lignes = texte if isinstance(texte, list) else [texte]
    for i, l in enumerate(lignes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = typo(texte_brut(l))
        r.font.size = PPt(taille)
        r.font.bold = gras
        r.font.italic = italique
        r.font.color.rgb = PRGB.from_string(couleur)
        p.space_after = PPt(taille * 0.6)
    return tb


def faire_pptx(r, chemin):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    W, Hh = prs.slide_width, prs.slide_height
    vide = prs.slide_layouts[6]
    m = 600000
    s = prs.slides.add_slide(vide)
    bg = s.shapes.add_shape(1, 0, 0, W, Hh)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRGB.from_string(INDIGO_L)
    bg.line.fill.background()
    boite(s, m, 1500000, W - 2 * m, 600000, r["kicker"].upper(), 14, INDIGO, True)
    boite(s, m, 2100000, W - 2 * m, 1800000, [r["t1"], r["t2"]], 44, NAVY, True)
    if r.get("sous"):
        boite(s, m, 4100000, W - 2 * m, 600000, r["sous"], 20, INDIGO, False, True)
    boite(s, m, Hh - 800000, W - 2 * m, 500000, "Modèle gratuit · Association ADéPA · les-extras.fr/ressources", 11, GRIS)
    for i, (court, grand, items) in enumerate(r["diapos"]):
        s = prs.slides.add_slide(vide)
        bande = s.shapes.add_shape(1, 0, 0, Emu(180000), Hh)
        bande.fill.solid()
        bande.fill.fore_color.rgb = PRGB.from_string(INDIGO)
        bande.line.fill.background()
        boite(s, m, 400000, W - 2 * m, 400000, f"{i + 1:02d} · {court}".upper(), 13, INDIGO, True)
        boite(s, m, 800000, W - 2 * m, 900000, grand, 34, NAVY, True)
        boite(s, m, 1900000, W - 2 * m, 3800000, ["À écrire : " + x for x in items], 20, GRIS, False, True)
        boite(s, m, Hh - 550000, W - 2 * m, 400000, f"{texte_brut(r['t1'])} {texte_brut(r['t2'])} · modèle ADéPA · {i + 1} / {len(r['diapos'])}", 10, GRIS)
    prs.core_properties.author = "Association ADéPA"
    prs.save(chemin)


def main():
    specs = json.load(open(sys.argv[1], encoding="utf-8"))
    out = sys.argv[2]
    seul = sys.argv[3] if len(sys.argv) > 3 else None
    os.makedirs(out, exist_ok=True)
    n = 0
    for r in specs:
        if seul and r["id"] != seul:
            continue
        for fmt in r.get("mod") or []:
            chemin = os.path.join(out, f"{r['id']}.{fmt}")
            {"docx": faire_docx, "xlsx": faire_xlsx, "pptx": faire_pptx}[fmt](r, chemin)
            n += 1
    print(n, "fichiers modifiables")


if __name__ == "__main__":
    main()
